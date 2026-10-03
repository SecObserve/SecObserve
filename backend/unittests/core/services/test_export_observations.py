from datetime import datetime, timezone
from io import BytesIO, StringIO
from unittest.mock import patch

from defusedcsv import csv
from django.http import HttpResponse
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from rest_framework.test import APIClient

from application.access_control.models import User
from application.commons.models import Settings
from application.core.models import (
    Branch,
    Observation,
    Product,
    Product_Member,
    Service,
)
from application.core.queries.observation import get_observations
from application.core.services.export_observations import (
    COLUMNS,
    EXCLUDED_FIELDS,
    export_observations_csv,
    export_observations_csv_for_product,
    export_observations_excel_for_product,
)
from application.import_observations.models import Parser
from application.rules.models import Rule
from unittests.base_test_case import BaseTestCase

HEADERS = [column.header for column in COLUMNS]
LAST_SEEN = datetime(2026, 9, 28, 10, 30, tzinfo=timezone.utc)


class TestExportObservations(BaseTestCase):
    def setUp(self):
        super().setUp()
        settings = Settings.load()
        settings.base_url_frontend = "https://secobserve.example.com"
        settings.save()

        parser = Parser.objects.create(name="parser_export", type="SCA", source="File")
        self.product_group = Product.objects.create(name="product_group_export", is_product_group=True)
        self.product_a = Product.objects.create(name="product_a_export", product_group=self.product_group)
        self.product_b = Product.objects.create(name="product_b_export", product_group=self.product_group)
        branch = Branch.objects.create(name="main", product=self.product_a)
        service = Service.objects.create(name="service_export", product=self.product_a)
        rule = Rule.objects.create(name="general_rule_export", parser=parser)

        self.observation_low = self._create(self.product_a, parser, "a low", "Low", branch=branch)
        self.observation_critical = self._create(
            self.product_a,
            parser,
            "=cmd|'/c calc'!A1",
            "Critical",
            branch=branch,
            origin_service=service,
            general_rule=rule,
            description="line 1\nline 2",
            assessment_vex_remediations=[{"category": "workaround", "text": "restart"}],
        )
        self.observation_other = self._create(self.product_b, parser, "b high", "High")

    def _create(self, product, parser, title, severity, **kwargs):
        return Observation.objects.create(
            product=product,
            parser=parser,
            title=title,
            current_severity=severity,
            parser_severity=severity,
            import_last_seen=LAST_SEEN,
            **kwargs,
        )

    def _excel_rows(self, workbook):
        buffer = BytesIO()
        workbook.save(buffer)
        worksheet = load_workbook(buffer).active
        return worksheet, [list(row) for row in worksheet.iter_rows(values_only=True)]

    def _csv_rows(self, response):
        return list(csv.reader(StringIO(response.content.decode("utf-8"))))

    def _value(self, row, header):
        return row[HEADERS.index(header)]

    def test_columns_cover_the_model(self):
        exported = {column.field.split("__")[0] for column in COLUMNS}
        annotations = {"observation_url", "product_url", "observation_log_comment"}
        fields = {field.name for field in Observation._meta.concrete_fields}

        self.assertEqual(set(), fields - exported - set(EXCLUDED_FIELDS))
        self.assertEqual(set(), exported - fields - annotations)
        self.assertEqual(len(HEADERS), len(set(HEADERS)))

    def test_excel_product_group(self):
        with self.assertNumQueries(2):
            workbook = export_observations_excel_for_product(self.product_group, None)
        worksheet, rows = self._excel_rows(workbook)

        self.assertEqual(HEADERS, rows[0])
        self.assertEqual(
            [self.observation_critical.pk, self.observation_low.pk, self.observation_other.pk],
            [row[0] for row in rows[1:]],
        )
        critical = rows[1]
        self.assertEqual(
            f"https://secobserve.example.com/#/observations/{self.observation_critical.pk}/show",
            self._value(critical, "Observation URL"),
        )
        self.assertEqual(
            f"https://secobserve.example.com/#/products/{self.product_a.pk}/show",
            self._value(critical, "Product URL"),
        )
        self.assertEqual("product_group_export", self._value(critical, "Product group"))
        self.assertEqual("product_a_export", self._value(critical, "Product"))
        self.assertEqual("main", self._value(critical, "Branch / Version"))
        self.assertEqual("service_export", self._value(critical, "Service"))
        self.assertEqual("parser_export", self._value(critical, "Parser"))
        self.assertEqual("general_rule_export", self._value(critical, "General fields rule"))
        self.assertEqual("'=cmd|'/c calc'!A1", self._value(critical, "Title"))
        self.assertEqual(
            '[{"category": "workaround", "text": "restart"}]', self._value(critical, "Assessment VEX remediations")
        )
        self.assertEqual(datetime(2026, 9, 28, 10, 30), self._value(critical, "Last seen"))
        self.assertIsNone(self._value(rows[3], "Branch / Version"))
        self.assertEqual("A2", worksheet.freeze_panes)
        self.assertEqual(f"A1:{get_column_letter(len(COLUMNS))}4", worksheet.auto_filter.ref)

    def test_excel_product_with_status(self):
        self.observation_low.current_status = "Resolved"
        self.observation_low.parser_status = "Resolved"
        self.observation_low.save()

        _, rows = self._excel_rows(export_observations_excel_for_product(self.product_a, ["Open"]))

        self.assertEqual([self.observation_critical.pk], [row[0] for row in rows[1:]])

    def test_excel_empty(self):
        _, rows = self._excel_rows(export_observations_excel_for_product(self.product_a, ["Duplicate"]))

        self.assertEqual([HEADERS], rows)

    def test_csv_product_group(self):
        response = HttpResponse(content_type="text/csv")
        with self.assertNumQueries(2):
            export_observations_csv_for_product(response, self.product_group, None)
        rows = self._csv_rows(response)

        self.assertEqual(HEADERS, rows[0])
        self.assertEqual(4, len(rows))
        critical = rows[1]
        self.assertEqual(str(self.observation_critical.pk), critical[0])
        self.assertEqual("line 1 NEWLINE line 2", self._value(critical, "Description"))
        self.assertEqual(
            '[{"category": "workaround", "text": "restart"}]', self._value(critical, "Assessment VEX remediations")
        )

    def test_csv_empty(self):
        response = HttpResponse(content_type="text/csv")
        export_observations_csv_for_product(response, self.product_a, ["Duplicate"])

        self.assertEqual([HEADERS], self._csv_rows(response))

    def test_csv_non_superuser(self):
        user = User.objects.create(username="user_export@example.com")
        Product_Member.objects.create(product=self.product_group, user=user, role=1)

        response = HttpResponse(content_type="text/csv")
        with patch("application.core.queries.observation.get_current_user", return_value=user):
            export_observations_csv(response, get_observations())
        rows = self._csv_rows(response)

        self.assertEqual(HEADERS, rows[0])
        self.assertEqual(4, len(rows))

    @patch("application.access_control.services.api_token_authentication.APITokenAuthentication.authenticate")
    def test_api_export_excel(self, mock_authenticate):
        mock_authenticate.return_value = User.objects.create(username="admin_export", is_superuser=True), None

        response = APIClient().get(f"/api/products/{self.product_group.pk}/export_observations_excel/")

        self.assertEqual(200, response.status_code)
        self.assertEqual('attachment; filename="observations.xlsx"', response["Content-Disposition"])
        worksheet = load_workbook(BytesIO(b"".join(response.streaming_content))).active
        self.assertEqual(4, worksheet.max_row)
